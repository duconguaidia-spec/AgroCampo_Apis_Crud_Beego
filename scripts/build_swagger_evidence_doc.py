import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(r"C:\Users\MAURICIO\OneDrive\Documentos\ADSO\AgroCampo\AgroCampo_Apis_Crud_Beego")
EVIDENCE_DIR = ROOT / "evidencias_swagger_escritorio_colores"
MANIFEST = EVIDENCE_DIR / "manifest.json"
OUT = ROOT / "Evidencias_Swagger_APIs_AgroCampo.docx"

METHOD_COLORS = {
    "GET": "61AFFE",
    "POST": "49CC90",
    "PUT": "FCA130",
    "DELETE": "F93E3E",
}


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, color="000000", size=9):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def style_doc(doc):
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)
    section.top_margin = Inches(0.35)
    section.bottom_margin = Inches(0.35)
    section.left_margin = Inches(0.45)
    section.right_margin = Inches(0.45)
    section.header_distance = Inches(0.2)
    section.footer_distance = Inches(0.2)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.05

    for name, size, color in [
        ("Title", 24, "0B2545"),
        ("Heading 1", 16, "2E74B5"),
        ("Heading 2", 13, "1F4D78"),
        ("Heading 3", 11, "1F4D78"),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(6)
        style.paragraph_format.space_after = Pt(4)


def add_cover(doc, entries):
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(100)
    run = title.add_run("Evidencias Swagger de APIs AgroCampo")
    run.bold = True
    run.font.size = Pt(28)
    run.font.color.rgb = RGBColor.from_string("0B2545")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(20)
    run = subtitle.add_run("Capturas en formato escritorio con colores por metodo")
    run.font.size = Pt(15)
    run.font.color.rgb = RGBColor.from_string("475569")

    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, text in enumerate(["API", "Operaciones", "Capturas", "Estados"]):
        set_cell_shading(hdr[i], "E8EEF5")
        set_cell_text(hdr[i], text, bold=True, color="0B2545", size=10)
    for api_name in sorted({e["apiName"] for e in entries}):
        rows = [e for e in entries if e["apiName"] == api_name]
        row = table.add_row().cells
        values = [
            api_name,
            str(len(rows)),
            str(len(rows) * 2),
            ", ".join(sorted({str(e["status"]) for e in rows})),
        ]
        for i, value in enumerate(values):
            set_cell_text(row[i], value, size=10)

    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.space_before = Pt(18)
    run = note.add_run("Cada operacion incluye dos pasos: 01_datos y 02_resultado.")
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor.from_string("475569")
    doc.add_page_break()


def add_index(doc, entries):
    doc.add_heading("Indice de Evidencias", level=1)
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    headers = ["API", "Recurso", "Metodo", "Endpoint", "Estado"]
    for i, text in enumerate(headers):
        set_cell_shading(table.rows[0].cells[i], "E8EEF5")
        set_cell_text(table.rows[0].cells[i], text, bold=True, color="0B2545", size=8)
    set_repeat_table_header(table.rows[0])

    for e in entries:
        row = table.add_row().cells
        values = [e["apiName"], e["resource"], e["method"], e["endpoint"], str(e["status"])]
        for i, value in enumerate(values):
            set_cell_text(row[i], value, size=7)
        set_cell_shading(row[2], METHOD_COLORS.get(e["method"], "DADCE0"))
        set_cell_text(row[2], e["method"], bold=True, color="FFFFFF", size=7)
    doc.add_page_break()


def add_section_header(doc, api_name):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(110)
    run = p.add_run(api_name)
    run.bold = True
    run.font.size = Pt(26)
    run.font.color.rgb = RGBColor.from_string("0B2545")
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p2.add_run("Evidencias organizadas por recurso, metodo y paso")
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor.from_string("64748B")
    doc.add_page_break()


def add_figure_page(doc, entry, img_path, step_label):
    color = METHOD_COLORS.get(entry["method"], "DADCE0")
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    cells = table.rows[0].cells
    metadata = [
        ("API", entry["apiName"]),
        ("Recurso", entry["resource"]),
        ("Metodo", entry["method"]),
        ("Paso", step_label),
    ]
    for idx, (label, value) in enumerate(metadata):
        set_cell_shading(cells[idx], color if label == "Metodo" else "E8EEF5")
        set_cell_text(cells[idx], f"{label}: {value}", bold=True, color="FFFFFF" if label == "Metodo" else "0B2545", size=8)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    r = p.add_run(f"{entry['method']} {entry['endpoint']}")
    r.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor.from_string("1F4D78")

    img_p = doc.add_paragraph()
    img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    img_p.paragraph_format.space_after = Pt(0)
    img_p.add_run().add_picture(str(img_path), width=Inches(9.65))
    doc.add_page_break()


def main():
    entries = json.loads(MANIFEST.read_text(encoding="utf-8"))
    entries.sort(key=lambda e: (e["apiName"], e["resource"].lower(), e["method"], e["endpoint"]))

    doc = Document()
    style_doc(doc)
    add_cover(doc, entries)
    add_index(doc, entries)

    current_api = None
    for entry in entries:
        if entry["apiName"] != current_api:
            current_api = entry["apiName"]
            add_section_header(doc, current_api)
        add_figure_page(doc, entry, Path(entry["input"]), "01_datos")
        add_figure_page(doc, entry, Path(entry["output"]), "02_resultado")

    # Remove trailing blank page introduced by the final page break.
    if doc.paragraphs and not doc.paragraphs[-1].text:
        p = doc.paragraphs[-1]._element
        p.getparent().remove(p)

    doc.core_properties.title = "Evidencias Swagger de APIs AgroCampo"
    doc.core_properties.subject = "Capturas Swagger organizadas por API, recurso, metodo y paso"
    doc.core_properties.author = "Codex"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
